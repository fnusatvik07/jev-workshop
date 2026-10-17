"""Small helper layer over python-pptx for the Jev workshop deck.

Palette and typography follow the house style for class decks:
cream background, ink text, coral accent, display serif for titles.
"""

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

CREAM = RGBColor(0xF0, 0xEE, 0xE6)
INK = RGBColor(0x1F, 0x1E, 0x1D)
CORAL = RGBColor(0xD9, 0x77, 0x57)
CLAY = RGBColor(0xBD, 0x5D, 0x3A)
SLATE = RGBColor(0x5E, 0x6B, 0x7A)
SAGE = RGBColor(0x7D, 0x8B, 0x6A)
MUTED = RGBColor(0x8A, 0x87, 0x80)
PAPER = RGBColor(0xFA, 0xF9, 0xF5)
LINE = RGBColor(0xDE, 0xDA, 0xCE)

SERIF = "Georgia"
SANS = "Helvetica Neue"
MONO = "Menlo"

W, H = Inches(13.333), Inches(7.5)


class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width = W
        self.prs.slide_height = H
        self.blank = self.prs.slide_layouts[6]
        self.n = 0

    # -- slide scaffolding -------------------------------------------------

    def slide(self, footer_text=None):
        s = self.prs.slides.add_slide(self.blank)
        self.n += 1
        bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
        bg.fill.solid()
        bg.fill.fore_color.rgb = CREAM
        bg.line.fill.background()
        bg.shadow.inherit = False
        if footer_text is not None:
            self.footer(s, footer_text)
        return s

    def footer(self, s, text):
        t = self.text(s, text, x=0.55, y=7.08, w=9.0, h=0.3, size=10, color=MUTED, font=SANS)
        self.text(s, str(self.n), x=12.5, y=7.08, w=0.4, h=0.3, size=10, color=MUTED, font=SANS, align=PP_ALIGN.RIGHT)
        return t

    def notes(self, s, text):
        s.notes_slide.notes_text_frame.text = text

    # -- primitives --------------------------------------------------------

    def rect(self, s, x, y, w, h, fill=PAPER, line=LINE, line_w=1.0, round_=True):
        shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if round_ else MSO_SHAPE.RECTANGLE
        r = s.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
        if round_:
            try:
                r.adjustments[0] = 0.055
            except Exception:
                pass
        if fill is None:
            r.fill.background()
        else:
            r.fill.solid()
            r.fill.fore_color.rgb = fill
        if line is None:
            r.line.fill.background()
        else:
            r.line.color.rgb = line
            r.line.width = Pt(line_w)
        r.shadow.inherit = False
        return r

    def text(self, s, content, x, y, w, h, size=18, color=INK, font=SANS, bold=False,
             italic=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line_spacing=1.0,
             space_after=0):
        """content: a string (newlines = paragraphs) or list of (text, overrides) runs per paragraph."""
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        paragraphs = content.split("\n") if isinstance(content, str) else content
        for i, para in enumerate(paragraphs):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            p.line_spacing = line_spacing
            if space_after:
                p.space_after = Pt(space_after)
            runs = [(para, {})] if isinstance(para, str) else para
            for run_text, ov in runs:
                r = p.add_run()
                r.text = run_text
                r.font.size = Pt(ov.get("size", size))
                r.font.name = ov.get("font", font)
                r.font.bold = ov.get("bold", bold)
                r.font.italic = ov.get("italic", italic)
                r.font.color.rgb = ov.get("color", color)
        return tb

    def title(self, s, text, size=33, y=0.52, x=0.55, w=12.2, color=INK):
        return self.text(s, text, x=x, y=y, w=w, h=0.95, size=size, color=color, font=SERIF)

    def kicker(self, s, text, y=0.22, x=0.57, color=CLAY):
        return self.text(s, text.upper(), x=x, y=y, w=10, h=0.3, size=12, color=color, font=SANS, bold=True)

    def chip(self, s, text, x, y, w=None, fill=CORAL, color=CREAM, size=12, h=0.34):
        w = w or (0.16 + 0.095 * len(text))
        r = self.rect(s, x, y, w, h, fill=fill, line=None)
        tf = r.text_frame
        tf.margin_left = tf.margin_right = Emu(0)
        tf.margin_top = tf.margin_bottom = Emu(0)
        tf.word_wrap = False
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = text
        run.font.size = Pt(size)
        run.font.bold = True
        run.font.name = SANS
        run.font.color.rgb = color
        return r

    def divider(self, s, x, y, w, color=LINE, weight=1.2):
        from pptx.enum.shapes import MSO_CONNECTOR
        ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y), Inches(x + w), Inches(y))
        ln.line.color.rgb = color
        ln.line.width = Pt(weight)
        ln.shadow.inherit = False
        return ln

    def bullets(self, s, items, x, y, w, size=16, gap=6, color=INK, line_spacing=1.08,
                marker="–  ", marker_color=None):
        paras = []
        for item in items:
            if isinstance(item, str):
                paras.append([(marker, {"color": marker_color or CORAL, "bold": True}), (item, {})])
            else:  # (lead, rest)
                lead, rest = item
                paras.append([
                    (marker, {"color": marker_color or CORAL, "bold": True}),
                    (lead, {"bold": True}),
                    (rest, {}),
                ])
        return self.text(s, paras, x=x, y=y, w=w, h=6.0, size=size, color=color,
                         line_spacing=line_spacing, space_after=gap)

    def code(self, s, code_text, x, y, w, h, size=12.5, title=None):
        box = self.rect(s, x, y, w, h, fill=RGBColor(0x2A, 0x28, 0x26), line=None)
        pad = 0.22
        if title:
            self.text(s, title, x=x + pad, y=y + 0.12, w=w - 2 * pad, h=0.28, size=10.5,
                      color=MUTED, font=MONO)
            code_y = y + 0.44
        else:
            code_y = y + pad
        self.text(s, code_text, x=x + pad, y=code_y, w=w - 2 * pad, h=h - (code_y - y) - pad,
                  size=size, color=RGBColor(0xED, 0xEA, 0xE3), font=MONO, line_spacing=1.12)
        return box

    def table(self, s, rows, x, y, w, col_widths, row_h=0.42, header=True, size=13,
              header_fill=INK, header_color=CREAM, zebra=True):
        n_rows, n_cols = len(rows), len(rows[0])
        tbl = s.shapes.add_table(n_rows, n_cols, Inches(x), Inches(y), Inches(w), Inches(row_h * n_rows)).table
        tbl.first_row = False
        tbl.horz_banding = False
        for j, cw in enumerate(col_widths):
            tbl.columns[j].width = Inches(cw)
        for i, row in enumerate(rows):
            tbl.rows[i].height = Inches(row_h)
            for j, val in enumerate(row):
                cell = tbl.cell(i, j)
                cell.margin_left = Inches(0.08)
                cell.margin_right = Inches(0.08)
                cell.margin_top = Inches(0.03)
                cell.margin_bottom = Inches(0.03)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                cell.fill.solid()
                if header and i == 0:
                    cell.fill.fore_color.rgb = header_fill
                elif zebra and i % 2 == 0:
                    cell.fill.fore_color.rgb = RGBColor(0xE9, 0xE6, 0xDC)
                else:
                    cell.fill.fore_color.rgb = PAPER
                tf = cell.text_frame
                tf.word_wrap = True
                text_val, ov = (val, {}) if isinstance(val, str) else val
                p = tf.paragraphs[0]
                run = p.add_run()
                run.text = text_val
                run.font.size = Pt(ov.get("size", size))
                run.font.name = ov.get("font", SANS)
                run.font.bold = ov.get("bold", header and i == 0)
                run.font.color.rgb = ov.get("color", header_color if (header and i == 0) else INK)
        return tbl

    def image(self, s, path, x, y, w=None, h=None):
        kw = {}
        if w:
            kw["width"] = Inches(w)
        if h:
            kw["height"] = Inches(h)
        return s.shapes.add_picture(path, Inches(x), Inches(y), **kw)

    def save(self, path):
        self.prs.save(path)
        return path
